"""One uploaded document turned into plain text, dispatched on extension. Table rows
become " | "-joined lines so a row survives chunking as one record."""

import csv
import logging
import os

from aws_resource_audit.errors import AuditError

logger = logging.getLogger(__name__)

_TEXT_EXTENSIONS = (".md", ".txt")


def extract_text(path: str) -> str:
    """The text of one document, or AuditError naming the file; rag.py skips that file."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in _TEXT_EXTENSIONS:
            return _read_text(path)
        if ext == ".csv":
            return _read_csv(path)
        if ext == ".pdf":
            return _read_pdf(path)
        if ext == ".docx":
            return _read_docx(path)
        if ext == ".xlsx":
            return _read_xlsx(path)
    except AuditError:
        raise
    except Exception as e:                                   # noqa: BLE001
        # The AuditError keeps only str(e); the original type and traceback are
        # the half that says WHICH parser gave up and where.
        logger.exception("parsing %s failed", path)
        raise AuditError(f"{os.path.basename(path)}: could not be read ({e})")
    raise AuditError(f"{os.path.basename(path)}: no reader for '{ext}' files.")


def _read_text(path: str) -> str:
    # errors="replace": a document that is mostly readable is worth indexing,
    # and a decode error here would otherwise drop the whole file.
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read()


def _read_csv(path: str) -> str:
    with open(path, encoding="utf-8", errors="replace", newline="") as f:
        return "\n".join(" | ".join(cell.strip() for cell in row)
                         for row in csv.reader(f) if any(c.strip() for c in row))


def _read_pdf(path: str) -> str:
    from pypdf import PdfReader

    pages = [page.extract_text() or "" for page in PdfReader(path).pages]
    text = "\n\n".join(p for p in pages if p.strip())
    if not text.strip():
        raise AuditError(
            f"{os.path.basename(path)}: no extractable text. A scanned PDF "
            f"is images of words, not words - run OCR on it first.")
    return text


def _read_docx(path: str) -> str:
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_xlsx(path: str) -> str:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    parts = []
    try:
        for sheet in book.worksheets:
            parts.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    parts.append(" | ".join(cells))
    finally:
        book.close()
    return "\n".join(parts)
